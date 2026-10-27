"""
Tests for Owner Collections feature and Excel (.xlsx) Exports.
Ensures accounting accuracy, zero duplicate revenue, pending confirmation workflows,
and valid Excel workbook generation across all reporting modules.
"""
import io
import openpyxl
import pytest
from datetime import date
from fastapi.testclient import TestClient

from app.main import app
from app.database import SessionLocal
from app.models.owners import Owner, OwnerCollection, OwnerCollectionStatus
from app.models.accounts import FinancialAccount
from app.services.journal import JournalService


@pytest.fixture
def client():
    return TestClient(app)


def test_record_confirmed_owner_collection(client):
    """Recording a confirmed collection with known amount updates correctly."""
    with SessionLocal() as db:
        cash_acct = db.query(FinancialAccount).filter(FinancialAccount.code == "1001").first()
        assert cash_acct is not None

    payload = {
        "collector_name": "Akash",
        "collection_date": date.today().isoformat(),
        "amount_paise": 600000,  # Rs. 6,000
        "destination_account_id": cash_acct.id,
        "source_category": "Pool Table Payments",
        "notes": "End of shift table cash collection",
        "is_unknown_amount": False,
    }

    res = client.post("/api/owners/collections", json=payload)
    assert res.status_code == 201
    data = res.json()
    assert data["reference"].startswith("COL-")
    assert data["collector_name"] == "Akash"
    assert data["amount_paise"] == 600000
    assert data["status"] == "CONFIRMED"
    # For Pool Table Payments kept in Cash in Hand (1001), no duplicate revenue entry is posted!
    assert data["journal_entry_id"] is None


def test_record_pending_owner_collection_and_confirm(client):
    """Recording an unknown amount leaves balances unchanged until confirmed."""
    with SessionLocal() as db:
        bank_acct = db.query(FinancialAccount).filter(FinancialAccount.code == "1002").first()
        initial_bank_bal = JournalService.get_account_balance(db, "1002")

    payload = {
        "collector_name": "Vishwa",
        "collection_date": date.today().isoformat(),
        "amount_paise": None,
        "destination_account_id": bank_acct.id,
        "source_category": "Other Business Income",
        "notes": "Pending envelope count from opening event",
        "is_unknown_amount": True,
    }

    res = client.post("/api/owners/collections", json=payload)
    assert res.status_code == 201
    data = res.json()
    coll_id = data["id"]
    assert data["status"] == "PENDING_CONFIRMATION"
    assert data["amount_paise"] is None
    assert data["journal_entry_id"] is None

    # Balance must remain completely unchanged!
    with SessionLocal() as db:
        assert JournalService.get_account_balance(db, "1002") == initial_bank_bal

    # Now confirm the collection
    confirm_payload = {
        "amount_paise": 1500000,  # Rs. 15,000
        "notes": "Counted and verified envelope contents",
        "confirmed_date": date.today().isoformat(),
    }
    c_res = client.post(f"/api/owners/collections/{coll_id}/confirm", json=confirm_payload)
    assert c_res.status_code == 200
    c_data = c_res.json()
    assert c_data["status"] == "CONFIRMED"
    assert c_data["amount_paise"] == 1500000
    assert c_data["journal_entry_id"] is not None

    # Balance should now reflect the confirmed income
    with SessionLocal() as db:
        new_bank_bal = JournalService.get_account_balance(db, "1002")
        assert new_bank_bal == initial_bank_bal + 1500000


def test_list_collections_with_filters(client):
    """Collections list endpoint supports filtering by partner, status, and dates."""
    res = client.get("/api/owners/collections?collector_name=Akash")
    assert res.status_code == 200
    items = res.json()
    assert isinstance(items, list)
    for it in items:
        assert "akash" in it["collector_name"].lower()


@pytest.mark.parametrize("export_endpoint,filename_keyword", [
    ("/api/reports/export/daily", "daily_report"),
    ("/api/reports/export/weekly", "weekly_report"),
    ("/api/reports/export/monthly?year=2026&month=10", "monthly_report"),
    ("/api/reports/export/ledger", "general_ledger"),
    ("/api/reports/export/transactions", "transactions"),
    ("/api/reports/export/expenses", "expenses"),
    ("/api/reports/export/cash-bank", "cash_and_bank"),
    ("/api/reports/export/savings", "business_savings"),
    ("/api/reports/export/partners?year=2026&month=10", "partners_distributions"),
    ("/api/reports/export/loans", "loans"),
    ("/api/reports/export/drinks", "soft_drinks"),
])
def test_excel_exports_produce_valid_workbooks(client, export_endpoint, filename_keyword):
    """Every Excel export endpoint returns 200 and a valid, readable .xlsx file."""
    res = client.get(export_endpoint)
    assert res.status_code == 200
    assert "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" in res.headers["content-type"]
    assert filename_keyword in res.headers["content-disposition"]

    # Load content into openpyxl to verify genuine, uncorrupted Excel structure
    buf = io.BytesIO(res.content)
    wb = openpyxl.load_workbook(buf)
    assert len(wb.sheetnames) >= 1
    ws = wb.active
    assert ws.max_row >= 4  # Must have title block and headers
