"""
Excel (.xlsx) Export Service for Pool Town ERP.
Produces genuine, beautifully-formatted Microsoft Excel workbooks using openpyxl.
Includes descriptive sheet names, professional headers, number/currency formatting,
column auto-sizing, and complete accounting details.
"""
from __future__ import annotations

import io
from datetime import date, datetime
from typing import Optional
from sqlalchemy.orm import Session
from sqlalchemy import func, desc, and_

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.sessions import PoolSession, PS4Session, SessionStatus
from app.models.payments import SessionPayment, PaymentStatus, PaymentMethod
from app.models.expenses import Expense, ExpenseStatus, ExpenseCategory
from app.models.inventory import Product, ProductCategory, ProductSale, StockPurchase, StockPurchaseItem
from app.models.accounts import FinancialAccount, AccountTransfer, CashCount, BankReconciliation
from app.models.journal import JournalEntry, JournalLine, JournalEntryStatus
from app.models.owners import (
    Owner, OwnerTransaction, OwnerTransactionType,
    Loan, LoanTransaction, LoanStatus, LoanTransactionType,
    ProfitAllocation, ProfitAllocationLine, OwnerCollection, OwnerCollectionStatus
)
from app.services.journal import JournalService
from app.routers.drinks import _get_stock_on_hand

# Palette matching Pool Town professional branding
NAVY_HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
SLATE_HEADER_FILL = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
ACCENT_BLUE_FILL = PatternFill(start_color="0284C7", end_color="0284C7", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

FONT_TITLE = Font(name="Calibri", size=15, bold=True, color="0F172A")
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="64748B")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_BOLD = Font(name="Calibri", size=11, bold=True, color="0F172A")
FONT_REGULAR = Font(name="Calibri", size=11, color="1E293B")
FONT_MUTED = Font(name="Calibri", size=10, color="64748B")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
CELL_BORDER = Border(left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE)
TOTAL_TOP_BORDER = Side(border_style="thin", color="0F172A")
TOTAL_BOTTOM_BORDER = Side(border_style="double", color="0F172A")
TOTAL_BORDER = Border(top=TOTAL_TOP_BORDER, bottom=TOTAL_BOTTOM_BORDER)

CURRENCY_FORMAT = "#,##0.00"
DATE_FORMAT = "YYYY-MM-DD"


def _auto_adjust_columns(ws, min_width=12):
    """Auto-adjusts column widths to fit content gracefully."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2):
                continue
            val_str = str(cell.value or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 3, min_width)


def _write_title_block(ws, title: str, subtitle: Optional[str] = None):
    """Writes standard title and subtitle block at top of worksheet."""
    ws.append([f"POOL TOWN — {title.upper()}"])
    ws.cell(row=1, column=1).font = FONT_TITLE
    if subtitle:
        ws.append([subtitle])
        ws.cell(row=2, column=1).font = FONT_SUBTITLE
    else:
        ws.append([f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}"])
        ws.cell(row=2, column=1).font = FONT_SUBTITLE
    ws.append([])  # blank row


def _apply_header_style(ws, row_idx: int, col_count: int, fill=NAVY_HEADER_FILL):
    """Applies bold white styling to table header row."""
    for c in range(1, col_count + 1):
        cell = ws.cell(row=row_idx, column=c)
        cell.font = FONT_HEADER
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center" if "date" in str(cell.value or "").lower() or "id" in str(cell.value or "").lower() else "left", vertical="center")
        cell.border = CELL_BORDER
    ws.row_dimensions[row_idx].height = 24


def export_daily_report(db: Session, report_date: date) -> io.BytesIO:
    """Generates comprehensive Daily Business & Accounting Report."""
    wb = openpyxl.Workbook()
    
    # ── Sheet 1: Daily Summary
    ws_sum = wb.active
    ws_sum.title = "Daily Summary"
    _write_title_block(ws_sum, "Daily Performance & Cash Report", f"Report Date: {report_date.isoformat()}")

    pool_rev = db.query(func.coalesce(func.sum(PoolSession.final_charge_paise), 0)).filter(
        PoolSession.session_date == report_date, PoolSession.status != SessionStatus.CANCELLED
    ).scalar() or 0
    ps4_rev = db.query(func.coalesce(func.sum(PS4Session.final_charge_paise), 0)).filter(
        PS4Session.session_date == report_date, PS4Session.status != SessionStatus.CANCELLED
    ).scalar() or 0
    drink_rev = db.query(func.coalesce(func.sum(ProductSale.total_revenue_paise), 0)).filter(
        ProductSale.sale_date == report_date
    ).scalar() or 0
    other_inc = db.query(func.coalesce(func.sum(JournalLine.amount_paise), 0)).join(JournalEntry).filter(
        JournalEntry.entry_date == report_date,
        JournalLine.account_code == "4004",
        JournalLine.is_debit == False,
        JournalEntry.status == JournalEntryStatus.POSTED
    ).scalar() or 0

    cash_coll = db.query(func.coalesce(func.sum(SessionPayment.amount_paise), 0)).filter(
        SessionPayment.collection_date == report_date,
        SessionPayment.payment_method == PaymentMethod.CASH,
        SessionPayment.status == PaymentStatus.CONFIRMED
    ).scalar() or 0
    bank_coll = db.query(func.coalesce(func.sum(SessionPayment.amount_paise), 0)).filter(
        SessionPayment.collection_date == report_date,
        SessionPayment.payment_method == PaymentMethod.BANK_TRANSFER,
        SessionPayment.status == PaymentStatus.CONFIRMED
    ).scalar() or 0

    expenses_paid = db.query(func.coalesce(func.sum(Expense.amount_paise), 0)).filter(
        Expense.expense_date == report_date, Expense.status != ExpenseStatus.CANCELLED
    ).scalar() or 0

    headers = ["Financial Category", "Amount (Rs.)", "Notes / Accounting Impact"]
    ws_sum.append(headers)
    _apply_header_style(ws_sum, 4, len(headers))

    rows = [
        ("Pool Table Revenue", pool_rev / 100.0, "Billiard table playtime billings"),
        ("PS4 Console Revenue", ps4_rev / 100.0, "Gaming lounge session billings"),
        ("Soft Drink Sales Revenue", drink_rev / 100.0, "Beverage counter sales"),
        ("Other Business Income", other_inc / 100.0, "Miscellaneous operating income"),
        ("TOTAL DAILY REVENUE", (pool_rev + ps4_rev + drink_rev + other_inc) / 100.0, "Gross business revenue earned"),
        ("Daily Cash Collected", cash_coll / 100.0, "Physical currency received"),
        ("Daily Bank Collected", bank_coll / 100.0, "Direct bank transfers received"),
        ("TOTAL COLLECTIONS", (cash_coll + bank_coll) / 100.0, "Total funds received"),
        ("Operating Expenses Paid", expenses_paid / 100.0, "Business payments and vendor expenses"),
        ("ESTIMATED NET DAILY CASH FLOW", ((cash_coll + bank_coll) - expenses_paid) / 100.0, "Net cash position change"),
    ]

    r_idx = 5
    for cat, amt, note in rows:
        ws_sum.append([cat, amt, note])
        c1 = ws_sum.cell(row=r_idx, column=1)
        c2 = ws_sum.cell(row=r_idx, column=2)
        c3 = ws_sum.cell(row=r_idx, column=3)
        c2.number_format = CURRENCY_FORMAT
        c1.border = CELL_BORDER
        c2.border = CELL_BORDER
        c3.border = CELL_BORDER
        if "TOTAL" in cat or "ESTIMATED" in cat:
            c1.font = FONT_BOLD
            c2.font = FONT_BOLD
            c2.border = TOTAL_BORDER
            c1.fill = ZEBRA_FILL
            c2.fill = ZEBRA_FILL
        r_idx += 1
    _auto_adjust_columns(ws_sum)

    # ── Sheet 2: Pool & PS4 Sessions
    ws_sess = wb.create_sheet(title="Sessions Detail")
    _write_title_block(ws_sess, "Sessions Detail", f"Sessions recorded on {report_date.isoformat()}")
    s_headers = ["Type", "Session Ref", "Player Name", "Recorded At", "Duration (mins)", "Total Charge (Rs.)", "Status"]
    ws_sess.append(s_headers)
    _apply_header_style(ws_sess, 4, len(s_headers), fill=SLATE_HEADER_FILL)

    pool_sessions = db.query(PoolSession).filter(PoolSession.session_date == report_date).order_by(desc(PoolSession.id)).all()
    ps4_sessions = db.query(PS4Session).filter(PS4Session.session_date == report_date).order_by(desc(PS4Session.id)).all()

    s_idx = 5
    for s in pool_sessions:
        ws_sess.append([
            "Pool Table", s.reference, s.player.name if s.player else "Walk-in",
            s.created_at.strftime("%H:%M") if s.created_at else "",
            s.actual_duration_minutes,
            (s.final_charge_paise or 0) / 100.0,
            s.status.value if hasattr(s.status, "value") else str(s.status)
        ])
        ws_sess.cell(row=s_idx, column=6).number_format = CURRENCY_FORMAT
        for c in range(1, 8):
            ws_sess.cell(row=s_idx, column=c).border = CELL_BORDER
        s_idx += 1

    for s in ps4_sessions:
        p_name = s.player1.name if hasattr(s, "player1") and s.player1 else "Walk-in"
        ws_sess.append([
            f"PS4 ({s.num_players} Players)", s.reference, p_name,
            s.created_at.strftime("%H:%M") if s.created_at else "",
            s.actual_duration_minutes,
            (s.final_charge_paise or 0) / 100.0,
            s.status.value if hasattr(s.status, "value") else str(s.status)
        ])
        ws_sess.cell(row=s_idx, column=6).number_format = CURRENCY_FORMAT
        for c in range(1, 8):
            ws_sess.cell(row=s_idx, column=c).border = CELL_BORDER
        s_idx += 1

    if s_idx == 5:
        ws_sess.append(["No sessions recorded for this day", "", "", "", "", 0, 0.0, ""])

    _auto_adjust_columns(ws_sess)

    # ── Sheet 3: Expenses Detail
    ws_exp = wb.create_sheet(title="Expenses Detail")
    _write_title_block(ws_exp, "Daily Expenses Detail", f"Expenses paid on {report_date.isoformat()}")
    e_headers = ["Reference", "Category", "Payee", "Description", "Amount (Rs.)", "Payment Method", "Status"]
    ws_exp.append(e_headers)
    _apply_header_style(ws_exp, 4, len(e_headers), fill=SLATE_HEADER_FILL)

    expenses = db.query(Expense).filter(Expense.expense_date == report_date).order_by(Expense.id).all()
    e_idx = 5
    for e in expenses:
        cat_name = e.category.name if e.category else "General Expense"
        ws_exp.append([
            e.reference, cat_name, e.payee or "N/A", e.description,
            e.amount_paise / 100.0,
            e.payment_method.value if hasattr(e.payment_method, "value") else str(e.payment_method),
            e.status.value if hasattr(e.status, "value") else str(e.status)
        ])
        ws_exp.cell(row=e_idx, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 8):
            ws_exp.cell(row=e_idx, column=c).border = CELL_BORDER
        e_idx += 1

    if e_idx == 5:
        ws_exp.append(["No expenses recorded for this day", "", "", "", 0.0, "", ""])

    _auto_adjust_columns(ws_exp)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_weekly_report(db: Session, week_date: date) -> io.BytesIO:
    """Generates Weekly Performance & Breakdown Report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Performance"

    from datetime import timedelta
    monday = week_date - timedelta(days=week_date.weekday())
    sunday = monday + timedelta(days=6)
    _write_title_block(ws, "Weekly Business Performance", f"Week: {monday.isoformat()} to {sunday.isoformat()}")

    headers = ["Date", "Day", "Pool Rev (Rs.)", "PS4 Rev (Rs.)", "Drinks Rev (Rs.)", "Other Inc (Rs.)", "Total Rev (Rs.)", "Cash Coll (Rs.)", "Bank Coll (Rs.)", "Expenses (Rs.)"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    curr = monday
    r_idx = 5
    tot_pool, tot_ps4, tot_drink, tot_other, tot_rev, tot_cash, tot_bank, tot_exp = 0, 0, 0, 0, 0, 0, 0, 0

    while curr <= sunday:
        p_rev = db.query(func.coalesce(func.sum(PoolSession.final_charge_paise), 0)).filter(
            PoolSession.session_date == curr, PoolSession.status != SessionStatus.CANCELLED
        ).scalar() or 0
        ps_rev = db.query(func.coalesce(func.sum(PS4Session.final_charge_paise), 0)).filter(
            PS4Session.session_date == curr, PS4Session.status != SessionStatus.CANCELLED
        ).scalar() or 0
        dr_rev = db.query(func.coalesce(func.sum(ProductSale.total_revenue_paise), 0)).filter(
            ProductSale.sale_date == curr
        ).scalar() or 0
        oth = db.query(func.coalesce(func.sum(JournalLine.amount_paise), 0)).join(JournalEntry).filter(
            JournalEntry.entry_date == curr,
            JournalLine.account_code == "4004",
            JournalLine.is_debit == False,
            JournalEntry.status == JournalEntryStatus.POSTED
        ).scalar() or 0
        t_rev = p_rev + ps_rev + dr_rev + oth

        c_coll = db.query(func.coalesce(func.sum(SessionPayment.amount_paise), 0)).filter(
            SessionPayment.collection_date == curr, SessionPayment.payment_method == PaymentMethod.CASH, SessionPayment.status == PaymentStatus.CONFIRMED
        ).scalar() or 0
        b_coll = db.query(func.coalesce(func.sum(SessionPayment.amount_paise), 0)).filter(
            SessionPayment.collection_date == curr, SessionPayment.payment_method == PaymentMethod.BANK_TRANSFER, SessionPayment.status == PaymentStatus.CONFIRMED
        ).scalar() or 0
        exp = db.query(func.coalesce(func.sum(Expense.amount_paise), 0)).filter(
            Expense.expense_date == curr, Expense.status != ExpenseStatus.CANCELLED
        ).scalar() or 0

        tot_pool += p_rev
        tot_ps4 += ps_rev
        tot_drink += dr_rev
        tot_other += oth
        tot_rev += t_rev
        tot_cash += c_coll
        tot_bank += b_coll
        tot_exp += exp

        ws.append([
            curr.strftime("%Y-%m-%d"), curr.strftime("%A"),
            p_rev / 100.0, ps_rev / 100.0, dr_rev / 100.0, oth / 100.0, t_rev / 100.0,
            c_coll / 100.0, b_coll / 100.0, exp / 100.0
        ])
        for c in range(3, 11):
            ws.cell(row=r_idx, column=c).number_format = CURRENCY_FORMAT
        for c in range(1, 11):
            ws.cell(row=r_idx, column=c).border = CELL_BORDER
        r_idx += 1
        curr += timedelta(days=1)

    # Weekly Total Row
    ws.append([
        "TOTAL", "Weekly Sum",
        tot_pool / 100.0, tot_ps4 / 100.0, tot_drink / 100.0, tot_other / 100.0, tot_rev / 100.0,
        tot_cash / 100.0, tot_bank / 100.0, tot_exp / 100.0
    ])
    for c in range(1, 11):
        cell = ws.cell(row=r_idx, column=c)
        cell.font = FONT_BOLD
        cell.border = TOTAL_BORDER
        cell.fill = ZEBRA_FILL
        if c >= 3:
            cell.number_format = CURRENCY_FORMAT

    _auto_adjust_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_monthly_report(db: Session, year: int, month: int) -> io.BytesIO:
    """Generates Monthly Financial & Partner Allocation Report."""
    from calendar import monthrange
    last_day = monthrange(year, month)[1]
    start_date = date(year, month, 1)
    end_date = date(year, month, last_day)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Monthly Executive Summary"
    _write_title_block(ws, "Monthly Accounting & Distribution Summary", f"Period: {year}-{month:02d} ({start_date} to {end_date})")

    p_rev = db.query(func.coalesce(func.sum(PoolSession.final_charge_paise), 0)).filter(
        PoolSession.session_date >= start_date, PoolSession.session_date <= end_date, PoolSession.status != SessionStatus.CANCELLED
    ).scalar() or 0
    ps_rev = db.query(func.coalesce(func.sum(PS4Session.final_charge_paise), 0)).filter(
        PS4Session.session_date >= start_date, PS4Session.session_date <= end_date, PS4Session.status != SessionStatus.CANCELLED
    ).scalar() or 0
    dr_rev = db.query(func.coalesce(func.sum(ProductSale.total_revenue_paise), 0)).filter(
        ProductSale.sale_date >= start_date, ProductSale.sale_date <= end_date
    ).scalar() or 0
    oth_rev = db.query(func.coalesce(func.sum(JournalLine.amount_paise), 0)).join(JournalEntry).filter(
        JournalEntry.entry_date >= start_date, JournalEntry.entry_date <= end_date,
        JournalLine.account_code == "4004", JournalLine.is_debit == False, JournalEntry.status == JournalEntryStatus.POSTED
    ).scalar() or 0
    gross_rev = p_rev + ps_rev + dr_rev + oth_rev

    cogs = db.query(func.coalesce(func.sum(JournalLine.amount_paise), 0)).join(JournalEntry).filter(
        JournalEntry.entry_date >= start_date, JournalEntry.entry_date <= end_date,
        JournalLine.account_code == "5001", JournalLine.is_debit == True, JournalEntry.status == JournalEntryStatus.POSTED
    ).scalar() or 0

    expenses = db.query(func.coalesce(func.sum(Expense.amount_paise), 0)).filter(
        Expense.expense_date >= start_date, Expense.expense_date <= end_date,
        Expense.status != ExpenseStatus.CANCELLED
    ).scalar() or 0

    net_profit = max(0, gross_rev - cogs - expenses)
    savings_share = int(net_profit * 0.40)
    akash_share = int(net_profit * 0.20)
    pasan_share = int(net_profit * 0.20)
    vishwa_share = net_profit - savings_share - akash_share - pasan_share

    headers = ["Metric / Description", "Amount (Rs.)", "Notes / Business Rule"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    summary_rows = [
        ("Pool Table Revenue", p_rev / 100.0, "Total billiards billings"),
        ("PS4 Gaming Revenue", ps_rev / 100.0, "Total PlayStation lounge billings"),
        ("Soft Drink Sales", dr_rev / 100.0, "Counter beverage sales"),
        ("Other Income", oth_rev / 100.0, "Miscellaneous operating income"),
        ("GROSS REVENUE", gross_rev / 100.0, "Total monthly business revenue"),
        ("Cost of Goods Sold (Beverages)", cogs / 100.0, "Direct drink purchase cost"),
        ("Operating Expenses", expenses / 100.0, "Rent, electricity, wages, maintenance, etc."),
        ("NET DISTRIBUTABLE PROFIT", net_profit / 100.0, "Eligible profit allocated according to partnership agreement"),
        ("Business Savings Allocation (40%)", savings_share / 100.0, "Transferred to business emergency & capital reserve"),
        ("Akash Monthly Share (20%)", akash_share / 100.0, "Allocated to partner Akash"),
        ("Pasan Monthly Share (20%)", pasan_share / 100.0, "Allocated to partner Pasan"),
        ("Vishwa Monthly Share (20%)", vishwa_share / 100.0, "Allocated to partner Vishwa"),
    ]

    r_idx = 5
    for m, val, note in summary_rows:
        ws.append([m, val, note])
        c1, c2, c3 = ws.cell(row=r_idx, column=1), ws.cell(row=r_idx, column=2), ws.cell(row=r_idx, column=3)
        c2.number_format = CURRENCY_FORMAT
        c1.border = CELL_BORDER
        c2.border = CELL_BORDER
        c3.border = CELL_BORDER
        if "GROSS" in m or "NET" in m or "Allocation" in m:
            c1.font = FONT_BOLD
            c2.font = FONT_BOLD
            if "NET" in m:
                c2.border = TOTAL_BORDER
                c1.fill = ZEBRA_FILL
                c2.fill = ZEBRA_FILL
        r_idx += 1
    _auto_adjust_columns(ws)

    # ── Sheet 2: Partner Settlement Detail
    ws_part = wb.create_sheet(title="Partner Settlement")
    _write_title_block(ws_part, "Partner Distribution Settlement", f"Month: {year}-{month:02d}")
    p_headers = ["Partner Name", "Profit Share %", "Monthly Share (Rs.)", "Advances Received (Rs.)", "Distributions Paid (Rs.)", "Remaining Payable (Rs.)"]
    ws_part.append(p_headers)
    _apply_header_style(ws_part, 4, len(p_headers), fill=SLATE_HEADER_FILL)

    owners = db.query(Owner).filter(Owner.is_active == True).order_by(Owner.id).all()
    p_idx = 5
    for o in owners:
        name_l = o.name.lower()
        if "akash" in name_l:
            share = akash_share
        elif "pasan" in name_l:
            share = pasan_share
        elif "vishwa" in name_l:
            share = vishwa_share
        else:
            share = int(net_profit * (o.profit_share_bps / 10000.0))

        adv = db.query(func.coalesce(func.sum(OwnerTransaction.amount_paise), 0)).filter(
            OwnerTransaction.owner_id == o.id,
            OwnerTransaction.transaction_type == OwnerTransactionType.ADVANCE,
            OwnerTransaction.period_year == year,
            OwnerTransaction.period_month == month
        ).scalar() or 0

        dist = db.query(func.coalesce(func.sum(OwnerTransaction.amount_paise), 0)).filter(
            OwnerTransaction.owner_id == o.id,
            OwnerTransaction.transaction_type.in_([OwnerTransactionType.DISTRIBUTION, OwnerTransactionType.WITHDRAWAL]),
            OwnerTransaction.period_year == year,
            OwnerTransaction.period_month == month
        ).scalar() or 0

        rem = max(0, share - adv - dist)

        ws_part.append([
            o.name, f"{o.profit_share_bps / 100.0:.1f}%",
            share / 100.0, adv / 100.0, dist / 100.0, rem / 100.0
        ])
        for c in range(3, 7):
            ws_part.cell(row=p_idx, column=c).number_format = CURRENCY_FORMAT
        for c in range(1, 7):
            ws_part.cell(row=p_idx, column=c).border = CELL_BORDER
        p_idx += 1

    _auto_adjust_columns(ws_part)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_general_ledger(db: Session, start_date: Optional[date] = None, end_date: Optional[date] = None) -> io.BytesIO:
    """Generates double-entry General Ledger audit export."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "General Ledger"
    
    sub = f"Date Filter: {start_date or 'Beginning'} to {end_date or 'Latest'}"
    _write_title_block(ws, "General Ledger & Audit Trail", sub)

    headers = ["Entry Date", "Journal Ref", "Account Code", "Account Name", "Debit (Rs.)", "Credit (Rs.)", "Description / Narration", "Source", "Status"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    query = db.query(JournalLine).join(JournalEntry).filter(JournalEntry.status == JournalEntryStatus.POSTED)
    if start_date:
        query = query.filter(JournalEntry.entry_date >= start_date)
    if end_date:
        query = query.filter(JournalEntry.entry_date <= end_date)

    lines = query.order_by(desc(JournalEntry.entry_date), desc(JournalEntry.id), JournalLine.id).limit(1000).all()

    tot_dr = 0
    tot_cr = 0
    r_idx = 5
    for l in lines:
        entry = l.entry
        dr_val = (l.amount_paise / 100.0) if l.is_debit else 0.0
        cr_val = (l.amount_paise / 100.0) if not l.is_debit else 0.0
        tot_dr += l.amount_paise if l.is_debit else 0
        tot_cr += l.amount_paise if not l.is_debit else 0

        ws.append([
            entry.entry_date.strftime("%Y-%m-%d"),
            entry.reference,
            l.account_code,
            l.account_name,
            dr_val,
            cr_val,
            entry.description or l.description,
            entry.source_type or "General",
            entry.status.value if hasattr(entry.status, "value") else str(entry.status)
        ])
        ws.cell(row=r_idx, column=5).number_format = CURRENCY_FORMAT
        ws.cell(row=r_idx, column=6).number_format = CURRENCY_FORMAT
        for c in range(1, 10):
            ws.cell(row=r_idx, column=c).border = CELL_BORDER
        r_idx += 1

    if r_idx == 5:
        ws.append(["No journal entries found for selected range", "", "", "", 0.0, 0.0, "", "", ""])
    else:
        # Total Row
        ws.append(["TOTAL", "", "", "", tot_dr / 100.0, tot_cr / 100.0, f"Balanced: {'YES' if tot_dr == tot_cr else 'NO'}", "", ""])
        for c in range(1, 10):
            cell = ws.cell(row=r_idx, column=c)
            cell.font = FONT_BOLD
            cell.border = TOTAL_BORDER
            cell.fill = ZEBRA_FILL
            if c in (5, 6):
                cell.number_format = CURRENCY_FORMAT

    _auto_adjust_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_transactions(db: Session, start_date: Optional[date] = None, end_date: Optional[date] = None, account_id: Optional[int] = None) -> io.BytesIO:
    """Generates Transaction History export."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Transaction History"
    
    sub = f"Filter: {start_date or 'All'} to {end_date or 'Latest'}"
    if account_id:
        acct = db.query(FinancialAccount).filter(FinancialAccount.id == account_id).first()
        if acct:
            sub += f" | Account: {acct.name}"

    _write_title_block(ws, "Transaction History", sub)

    headers = ["Date", "Journal Ref", "Account Name", "Debit (Rs.)", "Credit (Rs.)", "Narration", "Source Type"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    query = db.query(JournalLine).join(JournalEntry).filter(JournalEntry.status == JournalEntryStatus.POSTED)
    if start_date:
        query = query.filter(JournalEntry.entry_date >= start_date)
    if end_date:
        query = query.filter(JournalEntry.entry_date <= end_date)
    if account_id:
        acct = db.query(FinancialAccount).filter(FinancialAccount.id == account_id).first()
        if acct:
            query = query.filter(JournalLine.account_code == acct.code)

    lines = query.order_by(desc(JournalEntry.entry_date), desc(JournalEntry.id)).limit(1000).all()

    r_idx = 5
    for l in lines:
        entry = l.entry
        ws.append([
            entry.entry_date.strftime("%Y-%m-%d"),
            entry.reference,
            l.account_name,
            (l.amount_paise / 100.0) if l.is_debit else 0.0,
            (l.amount_paise / 100.0) if not l.is_debit else 0.0,
            entry.description or l.description,
            entry.source_type or "system"
        ])
        ws.cell(row=r_idx, column=4).number_format = CURRENCY_FORMAT
        ws.cell(row=r_idx, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 8):
            ws.cell(row=r_idx, column=c).border = CELL_BORDER
        r_idx += 1

    if r_idx == 5:
        ws.append(["No transactions found for selected criteria", "", "", 0.0, 0.0, "", ""])

    _auto_adjust_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_expenses(db: Session, start_date: Optional[date] = None, end_date: Optional[date] = None) -> io.BytesIO:
    """Generates Business Expenses report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses Ledger"

    sub = f"Period: {start_date or 'All Dates'} to {end_date or 'Latest'}"
    _write_title_block(ws, "Business Expenses Ledger", sub)

    headers = ["Date", "Expense Ref", "Category", "Payee / Vendor", "Description", "Amount (Rs.)", "Payment Method", "Status", "Notes"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    query = db.query(Expense)
    if start_date:
        query = query.filter(Expense.expense_date >= start_date)
    if end_date:
        query = query.filter(Expense.expense_date <= end_date)

    expenses = query.order_by(desc(Expense.expense_date), desc(Expense.id)).all()

    tot_amt = 0
    r_idx = 5
    for e in expenses:
        cat_name = e.category.name if e.category else "Other Expense"
        ws.append([
            e.expense_date.strftime("%Y-%m-%d"),
            e.reference,
            cat_name,
            e.payee or "N/A",
            e.description,
            e.amount_paise / 100.0,
            e.payment_method.value if hasattr(e.payment_method, "value") else str(e.payment_method),
            e.status.value if hasattr(e.status, "value") else str(e.status),
            e.notes or ""
        ])
        ws.cell(row=r_idx, column=6).number_format = CURRENCY_FORMAT
        for c in range(1, 10):
            ws.cell(row=r_idx, column=c).border = CELL_BORDER
        tot_amt += e.amount_paise
        r_idx += 1

    if r_idx == 5:
        ws.append(["No expenses recorded for selected criteria", "", "", "", "", 0.0, "", "", ""])
    else:
        ws.append(["TOTAL EXPENSES", "", "", "", "", tot_amt / 100.0, "", "", ""])
        for c in range(1, 10):
            cell = ws.cell(row=r_idx, column=c)
            cell.font = FONT_BOLD
            cell.border = TOTAL_BORDER
            cell.fill = ZEBRA_FILL
            if c == 6:
                cell.number_format = CURRENCY_FORMAT

    _auto_adjust_columns(ws)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_cash_and_bank(db: Session) -> io.BytesIO:
    """Generates Cash & Bank balances and movements export."""
    wb = openpyxl.Workbook()
    
    ws_bal = wb.active
    ws_bal.title = "Treasury Balances"
    _write_title_block(ws_bal, "Cash & Bank Balances Overview", "Current live ledger balances")

    headers = ["Account Classification", "Account Name", "Balance (Rs.)", "Operational Purpose"]
    ws_bal.append(headers)
    _apply_header_style(ws_bal, 4, len(headers))

    cash_bal = JournalService.get_account_balance(db, "1001")
    bank_bal = JournalService.get_account_balance(db, "1002")
    sav_bank = JournalService.get_account_balance(db, "1003")
    sav_cash = JournalService.get_account_balance(db, "1004")

    rows = [
        ("Operating Cash", "Cash in Hand", cash_bal / 100.0, "Physical currency available for daily operations"),
        ("Operating Bank", "Main Bank Account", bank_bal / 100.0, "Checking account for vendor bills & digital receipts"),
        ("SUBTOTAL OPERATING", "Total Available Operating Money", (cash_bal + bank_bal) / 100.0, "Immediately spendable business liquidity"),
        ("Business Savings", "Savings Held as Cash", sav_cash / 100.0, "Physical currency kept aside in emergency reserve"),
        ("Business Savings", "Savings Held in Bank", sav_bank / 100.0, "Bank deposit kept aside in business savings"),
        ("TOTAL SAVINGS", "Total Business Savings", (sav_cash + sav_bank) / 100.0, "Dedicated business savings fund (40% allocations)"),
        ("GRAND TOTAL", "Total Business Cash & Bank Assets", (cash_bal + bank_bal + sav_cash + sav_bank) / 100.0, "Total liquid wealth of the enterprise"),
    ]

    r_idx = 5
    for cat, name, amt, purpose in rows:
        ws_bal.append([cat, name, amt, purpose])
        c1, c2, c3, c4 = ws_bal.cell(row=r_idx, column=1), ws_bal.cell(row=r_idx, column=2), ws_bal.cell(row=r_idx, column=3), ws_bal.cell(row=r_idx, column=4)
        c3.number_format = CURRENCY_FORMAT
        c1.border = CELL_BORDER
        c2.border = CELL_BORDER
        c3.border = CELL_BORDER
        c4.border = CELL_BORDER
        if "TOTAL" in cat or "SUBTOTAL" in cat:
            c1.font = FONT_BOLD
            c2.font = FONT_BOLD
            c3.font = FONT_BOLD
            c3.border = TOTAL_BORDER
            c1.fill = ZEBRA_FILL
            c2.fill = ZEBRA_FILL
            c3.fill = ZEBRA_FILL
        r_idx += 1
    _auto_adjust_columns(ws_bal)

    # Sheet 2: Account Transfers
    ws_trf = wb.create_sheet(title="Transfers History")
    _write_title_block(ws_trf, "Account Transfers History", "Internal movements between business accounts")
    t_headers = ["Date", "Transfer Ref", "From Account", "To Account", "Amount (Rs.)", "Description / Reason"]
    ws_trf.append(t_headers)
    _apply_header_style(ws_trf, 4, len(t_headers), fill=SLATE_HEADER_FILL)

    transfers = db.query(AccountTransfer).order_by(desc(AccountTransfer.transfer_date), desc(AccountTransfer.id)).limit(200).all()
    t_idx = 5
    for t in transfers:
        ws_trf.append([
            t.transfer_date.strftime("%Y-%m-%d"),
            t.reference,
            t.from_account.name if t.from_account else "Unknown Account",
            t.to_account.name if t.to_account else "Unknown Account",
            t.amount_paise / 100.0,
            t.description or "Fund transfer"
        ])
        ws_trf.cell(row=t_idx, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 7):
            ws_trf.cell(row=t_idx, column=c).border = CELL_BORDER
        t_idx += 1

    if t_idx == 5:
        ws_trf.append(["No transfers recorded yet", "", "", "", 0.0, ""])

    _auto_adjust_columns(ws_trf)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_savings(db: Session) -> io.BytesIO:
    """Generates dedicated Business Savings report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Savings Reserve"
    _write_title_block(ws, "Business Savings & Reserve Fund", "Dedicated 40% net profit reserve")

    sav_cash = JournalService.get_account_balance(db, "1004")
    sav_bank = JournalService.get_account_balance(db, "1003")
    tot_sav = sav_cash + sav_bank

    headers = ["Savings Location", "Account Name", "Balance (Rs.)", "Status / Notes"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    ws.append(["Physical Currency", "Savings Held as Cash", sav_cash / 100.0, "Kept on premises in emergency safe"])
    ws.append(["Bank Account", "Savings Held in Bank", sav_bank / 100.0, "Kept in interest/savings bank account"])
    ws.append(["TOTAL SAVINGS", "Total Business Savings", tot_sav / 100.0, "Sum of physical and bank savings"])

    for r in range(5, 8):
        ws.cell(row=r, column=3).number_format = CURRENCY_FORMAT
        for c in range(1, 5):
            ws.cell(row=r, column=c).border = CELL_BORDER
    ws.cell(row=7, column=1).font = FONT_BOLD
    ws.cell(row=7, column=2).font = FONT_BOLD
    ws.cell(row=7, column=3).font = FONT_BOLD
    ws.cell(row=7, column=3).border = TOTAL_BORDER
    ws.cell(row=7, column=1).fill = ZEBRA_FILL
    ws.cell(row=7, column=2).fill = ZEBRA_FILL
    ws.cell(row=7, column=3).fill = ZEBRA_FILL

    _auto_adjust_columns(ws)

    # Sheet 2: Savings Movements
    ws_mov = wb.create_sheet(title="Savings Transfers")
    _write_title_block(ws_mov, "Savings Fund Movements", "Transfers to, from, or between savings accounts")
    m_headers = ["Date", "Reference", "Source Account", "Destination Account", "Amount (Rs.)", "Narration"]
    ws_mov.append(m_headers)
    _apply_header_style(ws_mov, 4, len(m_headers), fill=SLATE_HEADER_FILL)

    sav_acct_ids = [a.id for a in db.query(FinancialAccount).filter(FinancialAccount.code.in_(["1003", "1004"])).all()]
    transfers = db.query(AccountTransfer).filter(
        (AccountTransfer.from_account_id.in_(sav_acct_ids)) | (AccountTransfer.to_account_id.in_(sav_acct_ids))
    ).order_by(desc(AccountTransfer.transfer_date), desc(AccountTransfer.id)).all()

    m_idx = 5
    for t in transfers:
        ws_mov.append([
            t.transfer_date.strftime("%Y-%m-%d"),
            t.reference,
            t.from_account.name if t.from_account else "Unknown",
            t.to_account.name if t.to_account else "Unknown",
            t.amount_paise / 100.0,
            t.description or "Savings transfer"
        ])
        ws_mov.cell(row=m_idx, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 7):
            ws_mov.cell(row=m_idx, column=c).border = CELL_BORDER
        m_idx += 1

    if m_idx == 5:
        ws_mov.append(["No savings movements recorded yet", "", "", "", 0.0, ""])

    _auto_adjust_columns(ws_mov)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_partners(db: Session, year: Optional[int] = None, month: Optional[int] = None) -> io.BytesIO:
    """Generates Partners, Capital, Advances, and Collections report."""
    now = datetime.now()
    y = year or now.year
    m = month or now.month

    wb = openpyxl.Workbook()
    
    # Sheet 1: Monthly Settlement
    ws_set = wb.active
    ws_set.title = "Monthly Profit Sharing"
    _write_title_block(ws_set, "Partner Distribution & Advance Settlement", f"Period: {y}-{m:02d}")

    from calendar import monthrange
    last_d = monthrange(y, m)[1]
    st, en = date(y, m, 1), date(y, m, last_d)
    
    p_rev = db.query(func.coalesce(func.sum(PoolSession.final_charge_paise), 0)).filter(
        PoolSession.session_date >= st, PoolSession.session_date <= en, PoolSession.status != SessionStatus.CANCELLED
    ).scalar() or 0
    ps_rev = db.query(func.coalesce(func.sum(PS4Session.final_charge_paise), 0)).filter(
        PS4Session.session_date >= st, PS4Session.session_date <= en, PS4Session.status != SessionStatus.CANCELLED
    ).scalar() or 0
    dr_rev = db.query(func.coalesce(func.sum(ProductSale.total_revenue_paise), 0)).filter(
        ProductSale.sale_date >= st, ProductSale.sale_date <= en
    ).scalar() or 0
    oth_rev = db.query(func.coalesce(func.sum(JournalLine.amount_paise), 0)).join(JournalEntry).filter(
        JournalEntry.entry_date >= st, JournalEntry.entry_date <= en,
        JournalLine.account_code == "4004", JournalLine.is_debit == False, JournalEntry.status == JournalEntryStatus.POSTED
    ).scalar() or 0
    cogs = db.query(func.coalesce(func.sum(JournalLine.amount_paise), 0)).join(JournalEntry).filter(
        JournalEntry.entry_date >= st, JournalEntry.entry_date <= en,
        JournalLine.account_code == "5001", JournalLine.is_debit == True, JournalEntry.status == JournalEntryStatus.POSTED
    ).scalar() or 0
    exp = db.query(func.coalesce(func.sum(Expense.amount_paise), 0)).filter(
        Expense.expense_date >= st, Expense.expense_date <= en, Expense.status != ExpenseStatus.CANCELLED
    ).scalar() or 0

    net_prof = max(0, (p_rev + ps_rev + dr_rev + oth_rev) - cogs - exp)
    akash_sh = int(net_prof * 0.20)
    pasan_sh = int(net_prof * 0.20)
    vishwa_sh = net_prof - int(net_prof * 0.40) - akash_sh - pasan_sh

    headers = ["Partner Name", "Agreed Share %", "Monthly Share (Rs.)", "Advances Taken (Rs.)", "Distributions Paid (Rs.)", "Remaining Due (Rs.)"]
    ws_set.append(headers)
    _apply_header_style(ws_set, 4, len(headers))

    owners = db.query(Owner).filter(Owner.is_active == True).order_by(Owner.id).all()
    r_idx = 5
    for o in owners:
        name_l = o.name.lower()
        if "akash" in name_l:
            sh = akash_sh
        elif "pasan" in name_l:
            sh = pasan_sh
        elif "vishwa" in name_l:
            sh = vishwa_sh
        else:
            sh = int(net_prof * (o.profit_share_bps / 10000.0))

        adv = db.query(func.coalesce(func.sum(OwnerTransaction.amount_paise), 0)).filter(
            OwnerTransaction.owner_id == o.id,
            OwnerTransaction.transaction_type == OwnerTransactionType.ADVANCE,
            OwnerTransaction.period_year == y,
            OwnerTransaction.period_month == m
        ).scalar() or 0

        dist = db.query(func.coalesce(func.sum(OwnerTransaction.amount_paise), 0)).filter(
            OwnerTransaction.owner_id == o.id,
            OwnerTransaction.transaction_type.in_([OwnerTransactionType.DISTRIBUTION, OwnerTransactionType.WITHDRAWAL]),
            OwnerTransaction.period_year == y,
            OwnerTransaction.period_month == m
        ).scalar() or 0

        rem = max(0, sh - adv - dist)
        ws_set.append([o.name, f"{o.profit_share_bps / 100.0:.1f}%", sh / 100.0, adv / 100.0, dist / 100.0, rem / 100.0])
        for c in range(3, 7):
            ws_set.cell(row=r_idx, column=c).number_format = CURRENCY_FORMAT
        for c in range(1, 7):
            ws_set.cell(row=r_idx, column=c).border = CELL_BORDER
        r_idx += 1
    _auto_adjust_columns(ws_set)

    # Sheet 2: Collections by Partners
    ws_col = wb.create_sheet(title="Money Collected History")
    _write_title_block(ws_col, "Money Collected by Partners", "Records of all cash collected by Akash, Pasan, or Vishwa")
    c_headers = ["Date", "Collection Ref", "Collected By", "Amount (Rs.)", "Status", "Where Money Is Now", "Source Category", "Notes"]
    ws_col.append(c_headers)
    _apply_header_style(ws_col, 4, len(c_headers), fill=SLATE_HEADER_FILL)

    collections = db.query(OwnerCollection).order_by(desc(OwnerCollection.collection_date), desc(OwnerCollection.id)).all()
    c_idx = 5
    for c in collections:
        acct = db.query(FinancialAccount).filter(FinancialAccount.id == c.destination_account_id).first()
        amt_val = (c.amount_paise / 100.0) if c.amount_paise else "Pending Confirmation"
        ws_col.append([
            c.collection_date.strftime("%Y-%m-%d"),
            c.reference,
            c.collector_name,
            amt_val,
            c.status.value if hasattr(c.status, "value") else str(c.status),
            acct.name if acct else "Unknown",
            c.source_category,
            c.notes or ""
        ])
        if isinstance(amt_val, (int, float)):
            ws_col.cell(row=c_idx, column=4).number_format = CURRENCY_FORMAT
        for col_i in range(1, 9):
            ws_col.cell(row=c_idx, column=col_i).border = CELL_BORDER
        c_idx += 1

    if c_idx == 5:
        ws_col.append(["No partner collections recorded yet", "", "", "", "", "", "", ""])

    _auto_adjust_columns(ws_col)

    # Sheet 3: Capital & Advances Transactions
    ws_tx = wb.create_sheet(title="Partner Transactions")
    _write_title_block(ws_tx, "Partner Financial Events Ledger", "All capital contributions, advances, and distributions")
    tx_headers = ["Date", "Reference", "Partner", "Transaction Type", "Amount (Rs.)", "Payment Method / Account", "Period", "Description"]
    ws_tx.append(tx_headers)
    _apply_header_style(ws_tx, 4, len(tx_headers), fill=SLATE_HEADER_FILL)

    txs = db.query(OwnerTransaction).order_by(desc(OwnerTransaction.transaction_date), desc(OwnerTransaction.id)).all()
    tx_idx = 5
    for t in txs:
        o = db.query(Owner).filter(Owner.id == t.owner_id).first()
        acct = db.query(FinancialAccount).filter(FinancialAccount.id == t.account_id).first() if t.account_id else None
        ws_tx.append([
            t.transaction_date.strftime("%Y-%m-%d"),
            t.reference,
            o.name if o else "Unknown",
            t.transaction_type.value if hasattr(t.transaction_type, "value") else str(t.transaction_type),
            t.amount_paise / 100.0,
            acct.name if acct else "N/A",
            f"{t.period_year}-{t.period_month:02d}" if t.period_year and t.period_month else "N/A",
            t.description
        ])
        ws_tx.cell(row=tx_idx, column=5).number_format = CURRENCY_FORMAT
        for col_i in range(1, 9):
            ws_tx.cell(row=tx_idx, column=col_i).border = CELL_BORDER
        tx_idx += 1

    if tx_idx == 5:
        ws_tx.append(["No partner transactions recorded yet", "", "", "", 0.0, "", "", ""])

    _auto_adjust_columns(ws_tx)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_loans(db: Session) -> io.BytesIO:
    """Generates Business Loans & Repayments report."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Business Loans"
    _write_title_block(ws, "Business Loans & Repayments Schedule", "Track external, bank, and owner loans")

    headers = ["Loan Reference", "Lender Name", "Loan Type", "Original Principal (Rs.)", "Current Balance (Rs.)", "Start Date", "Interest Rate (%)", "Status", "Terms / Notes"]
    ws.append(headers)
    _apply_header_style(ws, 4, len(headers))

    loans = db.query(Loan).order_by(desc(Loan.start_date), desc(Loan.id)).all()
    r_idx = 5
    for loan in loans:
        repaid = db.query(func.coalesce(func.sum(LoanTransaction.amount_paise), 0)).filter(
            LoanTransaction.loan_id == loan.id,
            LoanTransaction.transaction_type == LoanTransactionType.PRINCIPAL_REPAYMENT
        ).scalar() or 0
        curr_bal = max(0, loan.original_principal_paise - repaid)

        ws.append([
            loan.reference,
            loan.lender_name,
            loan.loan_type.value if hasattr(loan.loan_type, "value") else str(loan.loan_type),
            loan.original_principal_paise / 100.0,
            curr_bal / 100.0,
            loan.start_date.strftime("%Y-%m-%d"),
            loan.interest_rate_bps / 100.0,
            loan.status.value if hasattr(loan.status, "value") else str(loan.status),
            loan.repayment_terms or loan.notes or ""
        ])
        ws.cell(row=r_idx, column=4).number_format = CURRENCY_FORMAT
        ws.cell(row=r_idx, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 10):
            ws.cell(row=r_idx, column=c).border = CELL_BORDER
        r_idx += 1

    if r_idx == 5:
        ws.append(["No loans recorded in the system", "", "", 0.0, 0.0, "", 0.0, "", ""])

    _auto_adjust_columns(ws)

    # Sheet 2: Repayments
    ws_rep = wb.create_sheet(title="Repayment Transactions")
    _write_title_block(ws_rep, "Loan Repayment Ledger", "Record of all principal and interest payments")
    r_headers = ["Date", "Transaction Ref", "Lender", "Transaction Type", "Amount Paid (Rs.)", "Payment Method Account", "Notes"]
    ws_rep.append(r_headers)
    _apply_header_style(ws_rep, 4, len(r_headers), fill=SLATE_HEADER_FILL)

    repayments = db.query(LoanTransaction).join(Loan).order_by(desc(LoanTransaction.transaction_date), desc(LoanTransaction.id)).all()
    rep_idx = 5
    for rep in repayments:
        acct = db.query(FinancialAccount).filter(FinancialAccount.id == rep.account_id).first() if rep.account_id else None
        ws_rep.append([
            rep.transaction_date.strftime("%Y-%m-%d"),
            rep.reference,
            rep.loan.lender_name if rep.loan else "Unknown",
            rep.transaction_type.value if hasattr(rep.transaction_type, "value") else str(rep.transaction_type),
            rep.amount_paise / 100.0,
            acct.name if acct else "N/A",
            rep.notes or ""
        ])
        ws_rep.cell(row=rep_idx, column=5).number_format = CURRENCY_FORMAT
        for c in range(1, 8):
            ws_rep.cell(row=rep_idx, column=c).border = CELL_BORDER
        rep_idx += 1

    if rep_idx == 5:
        ws_rep.append(["No loan repayments recorded yet", "", "", "", 0.0, "", ""])

    _auto_adjust_columns(ws_rep)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def export_drinks(db: Session, start_date: Optional[date] = None, end_date: Optional[date] = None) -> io.BytesIO:
    """Generates Soft Drinks inventory, valuation, and sales report."""
    wb = openpyxl.Workbook()
    
    # Sheet 1: Inventory & Stock Valuation
    ws_inv = wb.active
    ws_inv.title = "Current Stock Valuation"
    _write_title_block(ws_inv, "Soft Drinks Inventory & Valuation", "Current stock levels and retail pricing")

    headers = ["Product Name", "Category", "Stock on Hand", "Avg Cost (Rs.)", "Selling Price (Rs.)", "Total Inventory Cost (Rs.)", "Total Retail Value (Rs.)", "Reorder Level"]
    ws_inv.append(headers)
    _apply_header_style(ws_inv, 4, len(headers))

    products = db.query(Product).filter(Product.is_active == True).order_by(Product.name).all()
    r_idx = 5
    tot_cost_val = 0
    tot_retail_val = 0

    for p in products:
        soh = _get_stock_on_hand(db, p.id)
        cost_v = (soh * p.average_cost_paise) / 100.0
        ret_v = (soh * p.selling_price_paise) / 100.0
        tot_cost_val += (soh * p.average_cost_paise)
        tot_retail_val += (soh * p.selling_price_paise)

        ws_inv.append([
            p.name,
            p.category.name if p.category else "Drinks",
            soh,
            p.average_cost_paise / 100.0,
            p.selling_price_paise / 100.0,
            cost_v,
            ret_v,
            p.reorder_level
        ])
        for c in (4, 5, 6, 7):
            ws_inv.cell(row=r_idx, column=c).number_format = CURRENCY_FORMAT
        for c in range(1, 9):
            ws_inv.cell(row=r_idx, column=c).border = CELL_BORDER
        r_idx += 1

    if r_idx == 5:
        ws_inv.append(["No active products found", "", 0, 0.0, 0.0, 0.0, 0.0, 0])
    else:
        ws_inv.append(["TOTAL VALUATION", "", "", "", "", tot_cost_val / 100.0, tot_retail_val / 100.0, ""])
        for c in range(1, 9):
            cell = ws_inv.cell(row=r_idx, column=c)
            cell.font = FONT_BOLD
            cell.border = TOTAL_BORDER
            cell.fill = ZEBRA_FILL
            if c in (6, 7):
                cell.number_format = CURRENCY_FORMAT

    _auto_adjust_columns(ws_inv)

    # Sheet 2: Beverage Sales History
    ws_sales = wb.create_sheet(title="Beverage Sales")
    _write_title_block(ws_sales, "Soft Drink Sales History", f"Filter: {start_date or 'Beginning'} to {end_date or 'Latest'}")
    s_headers = ["Date", "Receipt Ref", "Product Name", "Quantity", "Selling Price (Rs.)", "Total Revenue (Rs.)", "Status"]
    ws_sales.append(s_headers)
    _apply_header_style(ws_sales, 4, len(s_headers), fill=SLATE_HEADER_FILL)

    query = db.query(ProductSale)
    if start_date:
        query = query.filter(ProductSale.sale_date >= start_date)
    if end_date:
        query = query.filter(ProductSale.sale_date <= end_date)

    sales = query.order_by(desc(ProductSale.sale_date), desc(ProductSale.id)).limit(1000).all()
    s_idx = 5
    tot_sales_paise = 0
    for s in sales:
        ws_sales.append([
            s.sale_date.strftime("%Y-%m-%d"),
            s.reference,
            s.product.name if s.product else "Drink",
            s.quantity,
            s.unit_price_paise / 100.0,
            s.total_revenue_paise / 100.0,
            "Recorded Sale"
        ])
        ws_sales.cell(row=s_idx, column=5).number_format = CURRENCY_FORMAT
        ws_sales.cell(row=s_idx, column=6).number_format = CURRENCY_FORMAT
        for c in range(1, 8):
            ws_sales.cell(row=s_idx, column=c).border = CELL_BORDER
        tot_sales_paise += s.total_revenue_paise
        s_idx += 1

    if s_idx == 5:
        ws_sales.append(["No sales found for selected range", "", "", 0, 0.0, 0.0, ""])
    else:
        ws_sales.append(["TOTAL SALES", "", "", "", "", tot_sales_paise / 100.0, ""])
        for c in range(1, 8):
            cell = ws_sales.cell(row=s_idx, column=c)
            cell.font = FONT_BOLD
            cell.border = TOTAL_BORDER
            cell.fill = ZEBRA_FILL
            if c == 6:
                cell.number_format = CURRENCY_FORMAT

    _auto_adjust_columns(ws_sales)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
